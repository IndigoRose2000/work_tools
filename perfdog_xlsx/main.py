from public import *
from openpyxl import load_workbook



def main():
    path = get_path()

    book1 = load_workbook(path[0])
    book2 = load_workbook(path[1])
    sheet1 = book1['all']
    sheet2 = book2['all']

    len_row1 = sheet1.max_row
    len_column1 = sheet1.max_column
    len_row2 = sheet2.max_row
    len_column2 = sheet2.max_column

    fir_xlsx = perf_data(sheet1, len_column1)
    sec_xlsx = perf_data(sheet2, len_column2)
    fir_cpu = app_cpu(sheet1, len_column1, len_row1)
    sec_cpu = app_cpu(sheet2, len_column2, len_row2)

    create_xlsx(path[2], path[3], fir_xlsx, sec_xlsx, fir_cpu, sec_cpu)



if __name__ == '__main__':
    main()