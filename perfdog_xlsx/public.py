from gooey import Gooey, GooeyParser
from openpyxl import Workbook
import re

@Gooey(program_name="get_perfdog_data_GUI", language="chinese")
def get_path():
    '''GUI工具获取对应信息'''
    parser = GooeyParser(description="自动提取数据")
    parser.add_argument('first_time', metavar="第一次录制", widget="TextField")
    parser.add_argument('secend_time', metavar="第二次录制", widget="TextField")
    # parser.add_argument('shebei', metavar="设备", choices=['Android', 'IOS'], widget="Dropdown")
    parser.add_argument('out_xlsx', metavar="输出xlsx名称（不用加.xlsx）", widget="TextField")
    args = parser.parse_args()  # 接收界面传递的参数

    pa = args.first_time.split('\\')
    path = ''
    for i in range(len(pa) - 1):
        path += pa[i] + '\\'
    path_01 = args.first_time
    path_02 = args.secend_time
    xlsx_out = args.out_xlsx
    # jiqi = args.shebei

    return [path_01, path_02, path, xlsx_out]


def perf_data(sheet, len_column):
    '''提取数据'''
    dic = {}
    for i in range(len_column):
        location1 = sheet.cell(8, i + 1).value
        location2 = sheet.cell(9, i + 1).value
        if location1 == location2 == None:
            continue
        if location1 == 'Avg(FPS)' or location1 == 'FPS>=25[%]'\
            or location1 == 'Avg(AppCPU)[%]' or location1 == 'AppCPU<=60%[%]'\
            or location1 == 'Avg(Memory)[MB]' or location1 == 'Peak(Memory)[MB]'\
            or location1 == '(Recv+Send)[KB/s]' or location1 == '(Recv+Send)[KB/10min]':
            dic[location1] = location2
    return dic

def app_cpu(sheet, len_column, len_row):
    '''整理数据'''
    ll = []
    for j in range(len_column):
        Top = sheet.cell(len_row - 2, j + 1).value
        Avg = sheet.cell(len_row - 1, j + 1).value
        Max = sheet.cell(len_row, j + 1).value

        if Top == Avg == Max == None:
            continue
        l = (Top, Avg, Max)
        if Top == 'AppCPU[%]':
            ll.append(l)
            return ll

def create_xlsx(path, xlsx_out, fir_xlsx, sec_xlsx, fir_cpu, sce_cpu):
    '''创建xlsx并写入数据'''
    wb = Workbook()
    sheet = wb.active

    top = ('Avg(FPS)', 'FPS>=25[%]', 'Avg(AppCPU)[%]', 'AppCPU<=60%[%]',
          '(Recv+Send)[KB/s]', '(Recv+Send)[KB/10min]', 'Avg(Memory)[MB]',
          'Peak(Memory)[MB]', 'AppCPU[%]')
    fir_perf = (fir_xlsx['Avg(FPS)'], fir_xlsx['FPS>=25[%]'], fir_xlsx['Avg(AppCPU)[%]'], fir_xlsx['AppCPU<=60%[%]'],
          fir_xlsx['(Recv+Send)[KB/s]'], fir_xlsx['(Recv+Send)[KB/10min]'], fir_xlsx['Avg(Memory)[MB]'], fir_xlsx['Peak(Memory)[MB]'],
          fir_cpu[0][2], fir_cpu[0][1])
    sec_perf = (sec_xlsx['Avg(FPS)'], sec_xlsx['FPS>=25[%]'], sec_xlsx['Avg(AppCPU)[%]'], sec_xlsx['AppCPU<=60%[%]'],
          sec_xlsx['(Recv+Send)[KB/s]'], sec_xlsx['(Recv+Send)[KB/10min]'], sec_xlsx['Avg(Memory)[MB]'], sec_xlsx['Peak(Memory)[MB]'],
          sce_cpu[0][2], sce_cpu[0][1])
    rows = (top, fir_perf, sec_perf)
    for row in rows:
        sheet.append(row)

    wb.save(f'{path}{xlsx_out}.xlsx')